#!/usr/bin/env python3
"""
Creates a clean, simple, and comprehensive Microsoft Word document (.docx) 
covering all 15 SATARK MPLADS Fraud & Anomaly Detectors.
Formatted in simple, natural document style: standard headings, clean typography, 
easy-to-read sections, plain-English explanations, formulas, thresholds, and real-world examples.
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner cell padding."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, fill_hex):
    """Set background color of a cell."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def style_simple_table(table, header_bg="2563EB"):
    """Styles a table cleanly with standard borders and clear headers."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header row
    for cell in table.rows[0].cells:
        set_cell_background(cell, header_bg)
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.name = "Calibri"
                r.font.size = Pt(10)
    
    # Body rows
    for r_idx, row in enumerate(table.rows[1:], start=1):
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for cell in row.cells:
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(30, 41, 59)

def generate_simple_docx(output_path: str):
    doc = Document()

    # Standard 1-inch margins
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # Base Normal Style
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(30, 41, 59) # Slate 800
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(5)

    # ------------------ TITLE & COVER ------------------
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("SATARK MPLADS Intelligence Portal")
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(16)
    r_sub = p_sub.add_run("Comprehensive Guide to the 15 Fraud & Anomaly Detectors")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    # Quick Metadata box
    doc.add_paragraph(
        "This document provides a simple, structured, and complete explanation of all 15 anomaly and fraud detection engines "
        "deployed in the SATARK MPLADS monitoring platform. For each detector, you will find its core purpose, data inputs, "
        "step-by-step logic, statistical thresholds, safeguards against false alarms, and a real-world example."
    )

    p_intro_rule = doc.add_paragraph()
    r_rule = p_intro_rule.add_run("Core Investigative Philosophy: ")
    r_rule.bold = True
    p_intro_rule.add_run(
        "Anomalies produced by these detectors are mathematical screening warnings and investigative leads. "
        "They highlight projects or entities requiring physical inspection or administrative verification. "
        "An alert does not automatically prove deliberate fraud."
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ------------------ SUMMARY TABLE ------------------
    h_sum = doc.add_heading(level=1)
    r_h_sum = h_sum.add_run("Master Summary of All 15 Detectors")
    r_h_sum.font.name = 'Arial'
    r_h_sum.font.size = Pt(15)
    r_h_sum.font.bold = True
    r_h_sum.font.color.rgb = RGBColor(30, 58, 138)

    sum_table = doc.add_table(rows=16, cols=4)
    headers = ["#", "Detector Name", "Plain-English Purpose", "Primary Trigger / Red Flag"]
    for c_idx, h in enumerate(headers):
        sum_table.cell(0, c_idx).paragraphs[0].text = h

    summary_rows = [
        ("01", "Unusual Pattern Detection", "Finds statistical outliers across cost and time", "Project is radically different from all peer works in district/category"),
        ("02", "Duplicate Work Detection", "Catches double-billing for the same physical work", "Nearly identical work descriptions within the same district"),
        ("03", "Cost Overrun Detection", "Checks if costs exceed official CPWD engineering rates", "Billed unit rate exceeds statutory ceiling by >= 5% and >= Rs. 10,000"),
        ("04", "Ghost Works Detection", "Identifies 'completed' projects with zero or tiny payments", "Project marked finished but ledger shows Rs. 0 or <50% paid"),
        ("05", "Bill Splitting (Smurfing)", "Catches splitting big projects to avoid tender rules", "Multiple works clustered just below Rs. 5 Lakh or Rs. 20 Lakh limits"),
        ("06", "Delays & Stalled Works", "Flags projects delayed far past the 1-year guideline", "Ongoing work exceeding 365 days or extreme historical completion delay"),
        ("07", "Suspicious Timing Forensics", "Detects year-end March dumping & pre-election rushes", "Over 40%–80% of budget spent in March, or sudden 2x+ pre-election speedup"),
        ("08", "Same-Day Bulk Completion", "Flags massive batches of works signed off on 1 day", "District daily completions exceed Mean + 3x StdDev (e.g. 20+ works)"),
        ("09", "Round-Number & Benford", "Catches fake or fabricated numbers in cost estimates", "High concentration of exact round numbers (Rs. 5L/Rs. 10L) violating Benford's Law"),
        ("10", "Vague Description Flag", "Flags projects with zero, short, or generic scopes", "Expensive projects described only as 'development work' or < 25 chars"),
        ("11", "Category-Amount Mismatch", "Ensures project cost is physically possible for work type", "Cost implausibly high (Rs. 28L for 1 handpump) or low (Rs. 45k for a school)"),
        ("12", "Verification Gap Flag", "Reconciles project sums with official MP ledger", "Sum of completed projects exceeds MP ledger balance by > 15%"),
        ("13", "IDA Risk Profiler", "Scores and ranks implementing districts on risk", "District-level Empirical Bayes composite risk score (Top 10% Critical)"),
        ("14", "MP & Constituency Profiler", "Scores and ranks MP portfolios on systemic risk", "Parliamentarian portfolio Empirical Bayes risk score (Top 10% Critical)"),
        ("15", "Copy-Paste Pricing", "Identifies reused boilerplate estimates across works", "Same MP has >= 5 works with identical cost across categories or identical unit rate")
    ]

    for r_idx, (d_num, d_name, d_desc, d_flag) in enumerate(summary_rows, start=1):
        sum_table.cell(r_idx, 0).paragraphs[0].text = d_num
        sum_table.cell(r_idx, 1).paragraphs[0].text = d_name
        sum_table.cell(r_idx, 2).paragraphs[0].text = d_desc
        sum_table.cell(r_idx, 3).paragraphs[0].text = d_flag

    style_simple_table(sum_table, header_bg="1E3A8A")
    doc.add_page_break()

    # ------------------ 15 DETAILED DETECTORS ------------------
    detectors = [
        {
            "id": "Detector 01",
            "name": "Unusual Pattern Detection (Isolation Forest)",
            "subtitle": "Multivariate Statistical Outlier Analysis",
            "plain_summary": (
                "Instead of looking at cost or duration by itself, this detector looks at the entire combination of features "
                "for every project simultaneously. It asks: 'Does this project look radically strange compared to typical "
                "projects in the same district and category?'"
            ),
            "inputs": [
                "Log-Transformed Project Cost: Normalizes wide expenditure differences (Rs. 2 Lakh to Rs. 5 Crore).",
                "Execution Duration (Days): Elapsed days from administrative recommendation to certified completion.",
                "Completion Month (1 to 12): Captures seasonal or fiscal timing patterns.",
                "Category Median Cost Delta: How far the project's cost deviates from the category median.",
                "District Median Cost Delta: How far the cost deviates from local district norms.",
                "Days per Lakh Spent: Efficiency metric (Duration ÷ Cost in Lakhs), spotting unusually rushed or stalled works."
            ],
            "rules": [
                "Minimum Dataset Guardrail: Requires at least 50 valid completed projects in the database. If fewer exist, the detector stops to prevent unreliable scores.",
                "Data Cleaning: Automatically removes projects with missing or invalid completion dates (e.g., completion before recommendation).",
                "Machine Learning Engine: Uses an Isolation Forest algorithm to partition multi-dimensional data points. Points that can be isolated with very few random cuts are marked as anomalies.",
                "Anomaly Scoring: Projects in the lowest score percentile are flagged."
            ],
            "severity": (
                "Severity scales smoothly from 0.00 to 1.00 based on how far outside the normal cluster the project falls. "
                "Marginal outliers receive lower severity (0.30–0.50), while extreme multi-dimensional outliers receive 0.80–1.00."
            ),
            "safeguards": (
                "A high score simply means the project is statistically unusual. Legitimate reasons include specialized foundation work, "
                "unforeseen rocky terrain, or emergency flood relief infrastructure."
            ),
            "workflow": [
                "Load all completed works with valid recommendation and completion dates.",
                "Calculate the 6 engineered features for every project.",
                "Verify that at least 50 completed projects exist.",
                "Fit the Isolation Forest model on the standardized feature matrix.",
                "Calculate the decision function anomaly score for each project.",
                "Filter against the system severity floor and save anomaly records with feature explanations."
            ],
            "example": {
                "desc": "Suppose a district typically builds community halls in 180 to 240 days for Rs. 12 to Rs. 15 Lakh.",
                "scenario": "A hall project is recorded as costing Rs. 48 Lakh and completed in only 9 days. While neither cost nor time alone might trigger basic filters, the combination of 3x normal cost in 5% of normal time is flagged immediately as a severe multi-dimensional outlier."
            }
        },
        {
            "id": "Detector 02",
            "name": "Duplicate Work Detection (NLP & Graph Clustering)",
            "subtitle": "Semantic Text Forensics & Union-Find Clustering",
            "plain_summary": (
                "Detects potential double-billing where two or more projects describe the exact same physical work "
                "under slightly different wording, abbreviations, or phrasing (e.g., 'Construction of PCC road from main gate to temple' "
                "vs. 'PCC road to mandir from entrance')."
            ),
            "inputs": [
                "Work Description Text: Natural language text describing the sanctioned scope of work.",
                "District Name: Used to partition comparisons (only works in the same district are compared).",
                "Work Category & MP Name: Used to confirm metadata agreement.",
                "Project Cost: Compared using cost ratio bounds (0.70x to 1.43x)."
            ],
            "rules": [
                "Length Gate: Descriptions with fewer than 20 characters (e.g., 'Road work') are excluded because they lack enough detail for text comparison.",
                "Template Exclusion: Common administrative phrases repeated more than 10 times across the database are ignored as boilerplate.",
                "Semantic Embeddings: Converts descriptions into high-dimensional numerical vectors using Sentence Transformers (or TF-IDF fallback).",
                "High-Confidence Duplicate: Flagged if Cosine Similarity >= 0.95, OR if Similarity >= 0.93 + identical Category + Cost within 70%–143%.",
                "Human Review Desk: Pairs with similarity between 0.88 and 0.93 with identical MP, Category, and Cost are routed to a pending verification queue.",
                "Union-Find Clustering: Groups connected duplicate pairs into clusters of 2 to 10 works."
            ],
            "severity": (
                "Base severity equals the cluster's average cosine similarity (e.g., 0.95 similarity = 0.95 severity). "
                "A +0.10 severity boost is added if all works in the cluster were recommended by the exact same MP."
            ),
            "safeguards": (
                "Cross-district comparisons are prohibited, eliminating false alarms between identical village names in different states. "
                "Phrases shorter than 20 characters and common boilerplate phrases are excluded."
            ),
            "workflow": [
                "Filter out descriptions shorter than 20 characters or repeated > 10 times.",
                "Generate vector embeddings for each description.",
                "Group embeddings strictly by district.",
                "Compute pairwise cosine similarity between all projects in each district.",
                "Apply matching rules (High-confidence >= 0.93/0.95; Pending review 0.88–0.93).",
                "Cluster matching pairs using Union-Find into groups of 2 to 10 works.",
                "Generate anomaly records for all works in the duplicate cluster."
            ],
            "example": {
                "desc": "Two project entries recorded under the same MP in District X:",
                "scenario": "Work 101: 'Construction of 200m cement concrete road near Shiva Mandir, Ward 4' (Rs. 8,00,000)\nWork 102: 'CC road construction 200m adjacent to Shiva temple Ward No 4' (Rs. 7,95,000)\nSimilarity score: 0.96. The detector flags both as duplicate cluster members."
            }
        },
        {
            "id": "Detector 03",
            "name": "Cost Overrun Detection (CPWD DSR Benchmarks)",
            "subtitle": "Statutory Central Public Works Department Rate Ceiling Enforcement",
            "plain_summary": (
                "Compares project costs against official statutory engineering rates established by the Central Public Works "
                "Department (CPWD Delhi Schedule of Rates 2023), adjusted for inflation and difficult geographic terrain."
            ),
            "inputs": [
                "CPWD DSR 2023 Master Rates: Baseline unit costs for roads, borewells, classrooms, street lights, halls, etc.",
                "Compounding Inflation: Automatic 6% annual escalation for completion years after 2023.",
                "Terrain Hardship Allowance: Additional +15% tolerance allowance for hilly and remote districts.",
                "Base Tolerance Buffer: +25% buffer above standard rates to allow for local market variance.",
                "Extracted Quantity: Linear meters, square meters, or item counts parsed from the description."
            ],
            "rules": [
                "Permissible Ceiling Formula: Maximum Allowed Rate = Base CPWD Rate x (1 + 0.06)^(Years post-2023) x (1 + 0.25 Base Buffer + 0.15 Terrain Buffer).",
                "Dual Overrun Condition: To be flagged, the actual unit rate must exceed the Permissible Ceiling by at least 5.0% AND the total excess spend must be at least Rs. 10,000.",
                "Extraction Confidence: High-confidence quantity parsing receives full severity; default or inferred quantities receive capped severity."
            ],
            "severity": (
                "Mapped directly to the percentage of overrun: 5% overrun (0.30 severity), 25% overrun (0.50), "
                "50% overrun (0.75), and >= 100% overrun (1.00)."
            ),
            "safeguards": (
                "Legitimate cost increases are fully protected by compounding inflation (6% per year), base tolerance (25%), "
                "and hilly terrain allowances (+15%). Minor overruns under Rs. 10,000 are ignored."
            ),
            "workflow": [
                "Map project category and description to CPWD benchmark master items.",
                "Calculate statutory ceiling including inflation and terrain allowances.",
                "Extract physical quantity from the text description.",
                "Compute actual billed unit rate: Total Cost ÷ Quantity.",
                "Check if actual unit rate exceeds ceiling by >= 5% and excess amount >= Rs. 10,000.",
                "Assign severity based on overrun percentage and persist evidence."
            ],
            "example": {
                "desc": "A plain-area district builds a 100m CC road in 2024.",
                "scenario": "CPWD 2023 benchmark: Rs. 3,500/m. With 6% inflation + 25% buffer, ceiling is Rs. 4,650/m. Billed cost is Rs. 9,00,000 for 100m (Rs. 9,000/m). Overrun is +93.5% with Rs. 4,35,000 excess expenditure -> Flagged with 0.95 Severity."
            }
        },
        {
            "id": "Detector 04",
            "name": "Ghost Works Detection (Phantom Projects)",
            "subtitle": "Ledger Reconciliation & Zero-Disbursement Screening",
            "plain_summary": (
                "Identifies projects certified as completed on official records but showing zero or near-zero financial payments "
                "in verified accounting ledgers, indicating paper-only projects or diversion of funds."
            ),
            "inputs": [
                "Project Status: Exclusively targets works with status = 'completed'.",
                "Verified Payment Ledger: Confirmed payment disbursements recorded in the transaction table.",
                "Disbursement Ratio: Total Amount Paid ÷ Total Sanctioned Cost.",
                "MP Portfolio Payment Gap: Overall unverified payment gap percentage across the MP's constituency."
            ],
            "rules": [
                "Signal 1 (Zero Payment): Confirmed payment ledger entry exists, work marked completed, but total paid is Rs. 0 (Base Severity: 0.80).",
                "Signal 2 (Severe Underpayment): Confirmed payment ledger entry exists, but total paid is less than 50% of claimed cost.",
                "Signal 3 (MP Portfolio Context): The sponsoring MP has an aggregate payment gap >= 40% across completed works.",
                "Compounding Rule: Final Severity = Highest Individual Signal Severity + 0.10 for each additional signal (max 1.00)."
            ],
            "severity": (
                "Base severity (0.60 to 0.80) adjusted by: (1) 30-day Recent Completion Grace: multiplied by 0.60 to allow for banking transit; "
                "(2) Small Projects (< Rs. 50,000): multiplied by 0.80."
            ),
            "safeguards": (
                "Missing payment records are treated as 'Unknown' rather than Rs. 0, preventing database synchronization delays from triggering false fraud alerts."
            ),
            "workflow": [
                "Filter for all projects with status == 'completed'.",
                "Verify whether an official payment transaction entry exists.",
                "Check for Rs. 0 paid or disbursement ratio < 50%.",
                "Check MP-level portfolio payment gap (>= 40%).",
                "Apply 30-day completion grace (0.60x) and small-work buffer (0.80x).",
                "Filter by severity floor and save anomaly record."
            ],
            "example": {
                "desc": "A community centre sanctioned for Rs. 25,00,000 is marked 'completed' 6 months ago.",
                "scenario": "The payment ledger shows exactly Rs. 0 released, and the MP portfolio has an overall 52% payment gap. Signal 1 (0.80) + Signal 3 (+0.10) triggers an anomaly with 0.90 Severity."
            }
        },
        {
            "id": "Detector 05",
            "name": "Bill Splitting / Smurfing (Procurement Fragmentation)",
            "subtitle": "Contract Threshold Evasion Forensics",
            "plain_summary": (
                "Catches deliberate splitting of large capital projects into multiple smaller contracts priced just below mandatory "
                "tender or administrative approval thresholds (specifically the Rs. 5 Lakh and Rs. 20 Lakh caps)."
            ),
            "inputs": [
                "Rs. 5 Lakh Band: Contracts priced between Rs. 4,50,000 and Rs. 4,99,999 (bypassing Rs. 5 Lakh open tender requirements).",
                "Rs. 20 Lakh Band: Contracts priced between Rs. 18,00,000 and Rs. 19,99,999 (bypassing Rs. 20 Lakh technical sanction caps).",
                "Temporal Grouping: Grouped by MP + Calendar Month + Threshold Band.",
                "Category Homogeneity: Evaluates whether fragmented contracts share the same category (e.g. all road works)."
            ],
            "rules": [
                "Rs. 5L Band Rule: 3 to 4 works in the same month by same MP -> Flagged (Base Severity 0.60); 5 or more works -> Flagged (Base Severity 0.80). Fewer than 3 works are ignored.",
                "Rs. 20L Band Rule: 2 or more works in the same month by same MP AND cumulative sum >= Rs. 20,00,000 -> Flagged (Base Severity 0.70).",
                "Category Boost: If all contracts in the cluster share the exact same category, add a +0.10 severity boost (max 1.00)."
            ],
            "severity": (
                "Base severity is 0.60 to 0.80 depending on cluster density, with an extra +0.10 boost for uniform categories."
            ),
            "safeguards": (
                "Isolated sub-5L contracts or works spread across different months are completely ignored. Smurfing alerts require both price proximity and temporal concentration."
            ),
            "workflow": [
                "Filter project costs into the 5L Band (Rs. 4.5L–<Rs. 5L) and 20L Band (Rs. 18L–<Rs. 20L).",
                "Group records by MP + Recommendation Month + Band.",
                "Apply density rules (>=3 works for 5L; >=2 works with sum >= Rs. 20L for 20L).",
                "Check for identical categories and apply +0.10 boost.",
                "Generate anomaly records for every constituent work in the cluster."
            ],
            "example": {
                "desc": "An MP sanctions 4 road works in October 2023 priced at Rs. 4,90,000, Rs. 4,95,000, Rs. 4,85,000, and Rs. 4,90,000.",
                "scenario": "Total expenditure is Rs. 19.6 Lakh. Instead of floating a single Rs. 20L tender, four contracts were issued just under the Rs. 5L tender limit. Cluster density >=3 + same category triggers 0.70 Severity."
            }
        },
        {
            "id": "Detector 06",
            "name": "Delays & Stalled Works (Timeline Forensics)",
            "subtitle": "Statutory 365-Day Completion Guideline Enforcement",
            "plain_summary": (
                "Identifies active public works that remain stalled long past the statutory 365-day (1-year) deadline, "
                "as well as finished works that suffered extreme, unapproved execution delays."
            ),
            "inputs": [
                "Recommendation & Completion Dates: Official administrative timestamps.",
                "Elapsed Operational Days: Calculated against the 365-day statutory completion guideline.",
                "Work Status: 'completed' vs active/ongoing.",
                "Multi-Phase Indicators: Detection of keywords ('phase', 'stage', 'part', 'package') signifying legitimate multi-year execution."
            ],
            "rules": [
                "Stalled In-Progress Works: Active works past 365 days. Severity scales: 365 days (0.50), 548 days / 1.5 yrs (0.65), 730 days / 2 yrs (0.80), 1095 days / 3+ yrs (1.00).",
                "Completed Delayed Works: Finished works where execution took > 365 days. Severity scales: 365 days (0.30), 548 days (0.45), 730 days (0.60), 1095 days (0.80).",
                "Multi-Phase Safeguard: Works with multi-phase keywords receive a 20% discount (0.80x) and are capped at 0.69 (preventing Critical tier classification)."
            ],
            "severity": (
                "Monotonically scaled from 0.30 to 1.00 based on elapsed delay past 1 year, discounted by 20% for multi-phase works."
            ),
            "safeguards": (
                "Differentiates active stalled works (higher risk) from finished delayed works (historical record). Multi-phase infrastructure packages are protected from emergency alarms."
            ),
            "workflow": [
                "Verify valid recommendation dates for all projects.",
                "Scan descriptions for multi-phase keywords ('phase', 'stage', 'package').",
                "Compute elapsed days for ongoing works OR duration for completed works.",
                "Flag works exceeding 365 days.",
                "Calculate severity using monotonic delay scales; apply 0.80x multiplier if multi-phase.",
                "Persist evidence including recommended date, days overdue, and status."
            ],
            "example": {
                "desc": "A drinking water pipeline sanctioned in January 2021 remains incomplete in October 2023 (1,000+ days).",
                "scenario": "The work is overdue by more than 635 days past the 365-day statutory guideline. Without multi-phase keywords, it receives a Critical 0.95 Severity alert for being stalled."
            }
        },
        {
            "id": "Detector 07",
            "name": "Suspicious Timing Forensics (March Dumping & Term Rush)",
            "subtitle": "Temporal Budget Dumping & Electoral Acceleration",
            "plain_summary": (
                "Identifies artificial project completions and expenditure surges driven by accounting deadlines "
                "(fiscal year-end March dumping) or political electoral cycles (pre-election term rush)."
            ),
            "inputs": [
                "Fiscal Calendar: Indian financial year (April 1 to March 31).",
                "March Concentration Index: Weighted composite: 40% March Works % + 60% March Expenditure %.",
                "Term-End Reference Date: Fixed statutory reference (May 31, 2024 for 17th Lok Sabha).",
                "Velocity Rush Ratio: Monthly completion rate during final 6 months divided by monthly rate during prior 54 months."
            ],
            "rules": [
                "March Fiscal Dumping: Evaluated per MP/FY. March Index mapped to severity: 30% (0.50), 45% (0.65), 65% (0.85), 85% (1.00). Normal expectation is 8.3% per month.",
                "Pre-Election Term Rush: Evaluated per MP (requires >= 5 works in baseline). Rush Ratio mapped to severity: 2.0x (0.50), 3.5x (0.65), 5.0x (0.80), 10.0x (1.00).",
                "Signal Combination: Final Severity = Max(March Dumping Severity, Term Rush Severity)."
            ],
            "severity": (
                "Takes the higher of the March Dumping or Term Rush severity scores. Filtered against the system severity floor."
            ),
            "safeguards": (
                "Expenditure has a 60% weight, so many minor projects completed in March do not trigger alarms unless large capital funds are dumped simultaneously."
            ),
            "workflow": [
                "Map project completion dates to Indian financial years.",
                "Calculate each MP's March project count %, March spending %, and weighted composite index.",
                "Calculate pre-election completion velocity vs. 54-month baseline.",
                "Determine severity for March dumping and term rush.",
                "Take the maximum severity and save anomaly evidence."
            ],
            "example": {
                "desc": "An MP completes 10 works per year from 2019 to 2023.",
                "scenario": "In the 6 months preceding the 2024 general election, 48 works are suddenly certified as completed (8 works/month vs. 0.83 works/month baseline). Rush Ratio = 9.6x -> Triggers Pre-Election Rush with 0.98 Severity."
            }
        },
        {
            "id": "Detector 08",
            "name": "Same-Day Bulk Completion (Daily Spike Forensics)",
            "subtitle": "Administrative Batch Sign-Off Screening",
            "plain_summary": (
                "Catches administrative batch approvals where an implausibly large number of distinct civil works are certified "
                "as completed on a single calendar day, indicating paperwork sign-offs without genuine on-site physical verification."
            ),
            "inputs": [
                "District Dynamic Baseline: Outlier-trimmed 95th percentile daily completion average and standard deviation.",
                "Spike Ratio: Total works completed on target date divided by normal daily district completion mean.",
                "March Year-End Dates: Special markers for March 25–30 (+0.10) and March 31 (+0.20 total boost).",
                "Category Diversity Ratio: Unique categories ÷ Total works on that date (< 0.20 indicates uniform batch sign-offs).",
                "MP Same-Day Count: Dedicated count of works closed by a single MP on that date."
            ],
            "rules": [
                "District Spike Threshold: Daily completions >= District Mean + 3x StdDev (minimum floor of 10 works).",
                "Spike Ratio Scale: 10x normal (0.40), 20x (0.60), 50x (0.85), 100x (1.00).",
                "Homogeneity Boost: Category Diversity < 0.20 adds +0.15 severity.",
                "MP Batch Threshold: 8 works on single day (0.50), 15 works (0.70), 25 works (0.90).",
                "Supply & Quarter Dampeners: Bulk supply keywords ('street light', 'lamp', 'led', 'pole', 'supply of') reduce severity by 30% (0.70x). Standard quarter-ends reduced by 20% (0.80x)."
            ],
            "severity": (
                "Max of district spike and MP batch severity, plus date and homogeneity boosts, discounted by supply and quarter-end dampeners."
            ),
            "safeguards": (
                "Bulk procurement contracts naturally delivered in batches (e.g., 50 solar street lights installed together) are protected by supply keyword filters."
            ),
            "workflow": [
                "Compute 95th percentile trimmed mean and StdDev of daily completions for each district.",
                "Aggregate completions by District + Date; flag dates exceeding Mean + 3x StdDev.",
                "Calculate Spike Ratio, March-end boosts, and category diversity ratio.",
                "Aggregate completions by MP + Date; flag single-day counts >= 8.",
                "Apply supply keyword (0.70x) and quarter-end (0.80x) discounts.",
                "Persist evidence including spike ratio, date, and constituent work IDs."
            ],
            "example": {
                "desc": "A district normally completes 1 to 2 civil works per week.",
                "scenario": "On March 31, exactly 42 separate school repairs and community halls are certified completed on that single day. Spike Ratio > 30x + March 31 boost triggers an anomaly with 0.95 Severity."
            }
        },
        {
            "id": "Detector 09",
            "name": "Round-Number Screen & Benford's Law Analysis",
            "subtitle": "Forensic Digit Distribution & Fabricated Budget Screening",
            "plain_summary": (
                "Checks whether project cost numbers follow natural statistical digit distributions (Benford's Law) "
                "or exhibit suspicious clustering at exact, artificial round numbers (such as exact multiples of Rs. 5 Lakh or Rs. 10 Lakh)."
            ),
            "inputs": [
                "First Significant Digit (1 to 9): Compared against Benford's theoretical distribution (log10(1 + 1/d)).",
                "Second Significant Digit (0 to 9): Evaluated when sample sizes are sufficiently large (>= 60 works).",
                "Roundness Level: Level 1 (Rs. 10k), Level 2 (Rs. 50k), Level 3 (Rs. 1L), Level 4 (Rs. 5L), Level 5 (Rs. 10L).",
                "Bonferroni Correction: Adjusts p-values across MP portfolios to control the false-discovery rate."
            ],
            "rules": [
                "Sample Size Guardrails: Requires at least 45 works for 1st-digit Chi-Square test, and 60 works for 2nd-digit test. Smaller portfolios are bypassed.",
                "Benford Violation: Adjusted p-value < 0.05 AND First-Digit Deviation > 0.15, OR Second-Digit adjusted p-value < 0.05.",
                "Round-Number Concentration: % of MP works with Level >= 3 (Rs. 1 Lakh+ multiples). Severity reference: 20% (0.50), 35% (0.70), 50% (0.90).",
                "Work-Level Flagging: A work is flagged if it is Level >= 4 (exact Rs. 5L or Rs. 10L) OR if the MP portfolio has >= 30% exact-lakh concentration."
            ],
            "severity": (
                "Base severity is 0.50, +0.20 for Level 5 (exact Rs. 10 Lakh multiples), compounded with MP portfolio roundness severity and Benford violations."
            ),
            "safeguards": (
                "An isolated round project is never flagged on its own. Flagging requires portfolio-level evidence that the MP repeatedly allocates artificial numbers."
            ),
            "workflow": [
                "Extract 1st/2nd digits and roundness levels for all non-zero project costs.",
                "Group by MP (>= 45 works) and run Chi-Square test against Benford distribution.",
                "Calculate % of MP works with roundness Level >= 3.",
                "Apply Bonferroni p-value corrections.",
                "Flag projects with Level >= 4 or inside >= 30% round portfolios.",
                "Save evidence with digits, Chi-Square statistics, and roundness percentages."
            ],
            "example": {
                "desc": "An MP portfolio contains 80 projects.",
                "scenario": "52 projects are budgeted at exactly Rs. 5,00,000 or Rs. 10,00,000 (65% round concentration). First-digit distribution completely fails Benford's Law (p < 0.001). Triggers an alert with 0.88 Severity."
            }
        },
        {
            "id": "Detector 10",
            "name": "Vague Description Flag (Text Forensics)",
            "subtitle": "Natural Language Engineering Detail & Specificity Scoring",
            "plain_summary": (
                "Flags project records with missing, excessively brief, generic boilerplate, or uninformative scopes of work. "
                "The core principle: significant public expenditure requires proportional engineering and location detail."
            ),
            "inputs": [
                "Character Length: Evaluated against cost-based length thresholds.",
                "Generic Phrases: Detection of uninformative terms ('development work', 'various works', 'miscellaneous', 'other work').",
                "5-Dimensional Specificity Score: Weighted metrics for Measurements (0.25), Locations (0.25), Technical Specs (0.20), Action Verb (0.15), and Beneficiaries (0.15).",
                "Corpus Template Repetition: Frequency of identical description strings across different works."
            ],
            "rules": [
                "Missing Description: Blank, 'None', 'Not specified', or 'nan' immediately receives Critical severity (1.00).",
                "Cost Exemption Floor: Projects under Rs. 2,00,000 are completely exempt from vagueness screening.",
                "High-Cost Works (>= Rs. 5 Lakh): Length < 25 chars -> Severity 0.85; 25–49 chars -> Severity 0.70.",
                "Medium-Cost Works (Rs. 2L–Rs. 5L): Length < 20 chars -> Severity 0.75; 20–39 chars -> Severity 0.60.",
                "Low Specificity Score: Score < 0.12 triggers low_specificity (0.80); Score < 0.20 for >= Rs. 5L triggers substandard_specificity (0.60).",
                "Template Reuse: Exact description repeated >= 10 times triggers template_repetition (0.50 to 0.85).",
                "Signal Compounding: Final Severity = Highest Signal + 0.15 per additional signal (capped at 1.00)."
            ],
            "severity": (
                "Combines length, generic phrasing, low specificity, and template reuse, with missing descriptions receiving automatic 1.00."
            ),
            "safeguards": (
                "Projects under Rs. 2 Lakh are exempted. Standard engineering terms (e.g., RCC, M20, chainage, village landmarks) directly boost specificity scores."
            ),
            "workflow": [
                "Check for null/blank descriptions (assign 1.00 severity).",
                "Skip projects with Cost < Rs. 2,00,000.",
                "Evaluate character length against Rs. 5L and Rs. 2L thresholds.",
                "Scan for generic phrases and evaluate 5-dimensional specificity score.",
                "Check corpus-wide template repetition (>= 10 matches).",
                "Combine signals (Max + 0.15 per additional) and persist evidence."
            ],
            "example": {
                "desc": "A Rs. 20,00,000 project is entered with the description: 'Development work'.",
                "scenario": "Character length is 16 chars (< 25 chars for >= Rs. 5L -> 0.85), contains generic phrase 'development work', and has specificity score 0.04 (< 0.12). Compounded severity = 1.00 Critical."
            }
        },
        {
            "id": "Detector 11",
            "name": "Category-Amount Mismatch (Engineering Plausibility)",
            "subtitle": "Physical Feasibility Ceilings & Unit Bounds",
            "plain_summary": (
                "Screens project budgets against physical engineering plausibility limits to catch expenditures that are either "
                "implausibly exorbitant (e.g. Rs. 28 Lakh for 1 handpump) or impossibly inadequate (e.g. Rs. 45,000 for an entire school building)."
            ),
            "inputs": [
                "Engineering Feasibility Master: Predefined bounds for 7 standard categories: School Building (Rs. 5L–Rs. 2Cr), Community Hall (Rs. 3L–Rs. 1.5Cr), Handpump (Rs. 25k–Rs. 1L), Overhead Tank (Rs. 4L–Rs. 30L), Compound Wall (Rs. 50k–Rs. 50L), Street Lights (Rs. 3k–Rs. 35k), CC Roads (Rs. 50k–Rs. 2.5Cr).",
                "Master File Override: Seamlessly updated via unit_prices_master.csv when available.",
                "Category Inference NLP: Maps generic classifications ('Others') to verified categories via keywords (e.g., 'borewell', 'LED light').",
                "Parsed Quantity: Physical quantity extracted from description (defaults to 1 unit if unstated)."
            ],
            "rules": [
                "Feasible Range Bounds: Effective Minimum = Unit Min x Quantity; Effective Maximum = Unit Max x Quantity.",
                "Implausibly Low Check: Actual Cost < Effective Minimum -> Flagged. Ratio = Min ÷ Cost. Severity: 1.0x (0.50), 5.0x (0.85), 10.0x (1.00).",
                "Implausibly High Check: Actual Cost > Effective Maximum -> Flagged. Ratio = Cost ÷ Max. Severity: 1.0x (0.50), 3.0x (0.80), 5.0x (1.00)."
            ],
            "severity": (
                "Calculated monotonically from the deviation ratio outside permissible engineering bounds, capped at 1.00."
            ),
            "safeguards": (
                "Projects lacking recognizable category keywords are safely skipped rather than evaluated against arbitrary numbers."
            ),
            "workflow": [
                "Load standard unit bounds (with CSV master overrides if available).",
                "Infer normalized category from text and extract quantity (default 1).",
                "Compute Effective Minimum and Maximum bounds.",
                "Test if Cost < Minimum (implausibly low) OR Cost > Maximum (implausibly high).",
                "Map deviation ratio to severity and save anomaly record."
            ],
            "example": {
                "desc": "Category: Drinking Water - Handpump; Quantity: 1; Recorded Cost: Rs. 28,00,000.",
                "scenario": "Maximum feasible engineering cost for a handpump is Rs. 1,00,000. Ratio = Rs. 28L ÷ Rs. 1L = 28.0x. Triggers Implausibly High alert with maximum 1.00 Severity."
            }
        },
        {
            "id": "Detector 12",
            "name": "Verification Gap Flag (Ledger Reconciliation)",
            "subtitle": "Macro Ledger vs. Micro Project Reconciliation",
            "plain_summary": (
                "Reconciles bottom-up completed project expenditure against top-down MP-level official financial ledgers, "
                "spotting discrepancies where claimed completed projects exceed official ledger balances, or where completed projects have severe disbursement deficits."
            ),
            "inputs": [
                "Official Financial Ledger: Government financial breakdown schedule (all_mps_financial_breakdown.csv).",
                "Bottom-Up Project Sum: Sum of all completed project costs in the database for each MP.",
                "Divergence Ratio: Bottom-Up Project Sum ÷ Official Ledger Completed Value.",
                "Work Disbursement Ratio: Verified paid amount ÷ Project cost evaluated against MP portfolio payment gaps."
            ],
            "rules": [
                "Macro Ledger Divergence: Triggered when Database Sum exceeds Official Ledger Completed Value by > 15% (Divergence Ratio > 1.15). Severity: 1.15x (0.50), 1.50x (0.70), 2.50x (1.00). Attributed to all constituent works.",
                "Micro Disbursement Deficit: Triggered when: Verified payment record exists AND Disbursement Ratio < 25% AND MP Portfolio Payment Gap >= 60%. Severity scales with (1 - Disbursement Ratio).",
                "Composite Selection: Final Severity = Max(Macro Divergence Severity, Micro Deficit Severity)."
            ],
            "severity": (
                "Takes the higher of the macro ledger divergence or micro disbursement deficit severity. Filtered against the severity floor."
            ),
            "safeguards": (
                "Missing payment records are never treated as Rs. 0 paid, preventing database pipeline delays from being misidentified as accounting fraud."
            ),
            "workflow": [
                "Load MP official financial breakdown schedule.",
                "Sum costs of completed works per MP in the database.",
                "If Database Sum ÷ Ledger Value > 1.15, compute macro divergence severity.",
                "Audit individual completed works for payment existence, disbursement ratio (< 25%), and MP gap (>= 60%).",
                "Take Max(Divergence, Deficit) severity and persist evidence."
            ],
            "example": {
                "desc": "An MP's database entries show Rs. 14.5 Crore of completed projects.",
                "scenario": "The official verified financial ledger records only Rs. 8.0 Crore of completed works. Divergence Ratio = 1.81x (> 1.15x threshold) -> Triggers Verification Gap with 0.81 Severity across constituent works."
            }
        },
        {
            "id": "Detector 13",
            "name": "Implementing District Authority (IDA) Risk Profiler",
            "subtitle": "Hierarchical Empirical Bayes Entity Risk Aggregator",
            "plain_summary": (
                "Aggregates anomaly signals across all detectors to compute a single, statistically balanced composite risk score "
                "and relative percentile ranking for each Implementing District Authority (IDA)."
            ),
            "inputs": [
                "Cross-Detector Anomalies: Outputs from Detectors 1 through 12 for the active run.",
                "Unique Flagged Project Rate: Unique flagged projects ÷ Total district projects (avoids double-counting).",
                "Weighted Detector Violations: Incorporates detector-specific violation frequencies and configured weights.",
                "Empirical Bayes Shrinkage (m=30): Adjusts district scores toward the national average to stabilize small-sample districts."
            ],
            "rules": [
                "Raw Risk Score Formula: Sum of (Violation Rate x 100 x Mean Severity x Detector Weight) across all detectors, capped at 100.",
                "Empirical Bayes Shrinkage: Shrunk Risk = Raw Risk x [N / (N + 30)] + National Mean x [30 / (N + 30)], where N is total district projects.",
                "Percentile Risk Tiers: Top 10% of districts by shrunk risk -> Critical; Next 20% (10%–30%) -> High; Next 30% (30%–60%) -> Medium; Remaining 40% -> Clean."
            ],
            "severity": (
                "A continuous composite score (0–100) representing shrunk systemic risk, converted into relative administrative tiers."
            ),
            "safeguards": (
                "Empirical Bayes shrinkage prevents small districts with only 5–10 projects from artificially dominating the Critical tier due to high raw percentages."
            ),
            "workflow": [
                "Join projects and active run anomalies by district.",
                "Calculate total works, expenditure, and unique flagged works per district.",
                "Compute detector-specific violation rates and weighted contributions.",
                "Calculate raw risk score (capped at 100).",
                "Apply Empirical Bayes shrinkage formula (m=30).",
                "Rank districts descending by shrunk risk and assign percentile tiers (Critical/High/Medium/Clean).",
                "Save EntityRisk record with detector breakdown."
            ],
            "example": {
                "desc": "District A has 200 projects, 45 of which trigger various detectors (overruns, ghost works, delays).",
                "scenario": "Its weighted raw score is 68.2. Because N=200 is large, shrinkage is minimal (weight 200/230 = 87% on raw score). Shrunk risk = 64.5, placing District A in the Top 10% Critical tier for targeted audit."
            }
        },
        {
            "id": "Detector 14",
            "name": "MP & Constituency Risk Profiler",
            "subtitle": "Hierarchical Empirical Bayes Portfolio Risk Aggregator",
            "plain_summary": (
                "Synthesizes anomaly results across an entire parliamentarian's project portfolio, evaluating the breadth, frequency, "
                "and severity of irregular patterns to establish an executive risk profile and comparative ranking."
            ),
            "inputs": [
                "Constituency & MP Binding: Maps all recommended works to the sponsoring MP and constituency.",
                "Multi-Signal Portfolio Profiling: Integrates duplicate works, cost overruns, smurfing, delays, Benford violations, and timing anomalies.",
                "Weighted Portfolio Score: Incorporates detector-specific violation frequencies and configured weights.",
                "Empirical Bayes Shrinkage (m=20): Stabilizes risk scores for parliamentarians with small numbers of recommended works."
            ],
            "rules": [
                "Raw Portfolio Score: Sum of (Violation Rate x 100 x Mean Severity x Detector Weight), capped at 100.",
                "Empirical Bayes Formula: Shrunk Risk = Raw Risk x [N / (N + 20)] + National MP Mean x [20 / (N + 20)], where N is total MP works.",
                "Percentile Risk Tiers: Top 10% of MPs by shrunk risk -> Critical; Next 20% (10%–30%) -> High; Next 30% (30%–60%) -> Medium; Remaining 40% -> Clean."
            ],
            "severity": (
                "Composite score (0–100) scaled via Empirical Bayes shrinkage, mapped to relative percentile tiers."
            ),
            "safeguards": (
                "Small portfolios (e.g. newly elected MPs with under 10 works) are stabilized by m=20 shrinkage so that a single delayed work doesn't produce an artificial Critical rating."
            ),
            "workflow": [
                "Group project and anomaly records by MP name.",
                "Calculate total works, expenditure, and unique flagged works per MP.",
                "Calculate violation rates and weighted scores per detector.",
                "Sum detector contributions into raw score (capped at 100).",
                "Apply Empirical Bayes shrinkage formula (m=20).",
                "Rank MPs descending by shrunk risk and assign percentile tiers.",
                "Persist EntityRisk record with complete detector breakdown JSON."
            ],
            "example": {
                "desc": "An MP recommends 120 projects. 35 are flagged across smurfing, vague descriptions, and round-number screen.",
                "scenario": "Weighted raw risk = 58.4. Shrunk score = 54.1, placing this MP in the 88th percentile (High Risk tier), warranting portfolio-level review."
            }
        },
        {
            "id": "Detector 15",
            "name": "Copy-Paste Pricing (Cloned Estimates Forensics)",
            "subtitle": "Combinatorial Pricing & Template Replication Analysis",
            "plain_summary": (
                "Identifies cloned or copy-pasted budgeting where unrelated projects under the same MP share identical total project costs "
                "across different categories, or share identical rounded unit rates within the same category."
            ),
            "inputs": [
                "Cross-Category Cost Clustering: Groups works by MP + Exact Rounded Cost across multiple distinct categories.",
                "High-Confidence Unit Rate Extraction: Extracts quantity and calculates unit rate (Cost ÷ Quantity).",
                "Unit Rate Rounding: Rounds calculated unit rates to the nearest Rs. 100.",
                "Category-Specific Unit Rate Clustering: Evaluates repeated unit pricing within the same MP and category."
            ],
            "rules": [
                "Signal 1 (Cross-Category Exact Cost Clone): Triggered when >= 5 works under the same MP share the exact same cost AND span > 1 category. Severity scale: 5 works (0.75), 10 works (0.88), 20 works (0.98).",
                "Signal 2 (Cloned Unit Rate): Triggered when >= 5 works under the same MP and category share the identical rounded unit rate (to nearest Rs. 100) with high extraction confidence. Severity scale: 5 repeats (0.50), 10 repeats (0.70), 20 repeats (0.90).",
                "Signal Combination: Final Severity = Max(Cross-Category Severity, Cloned Unit Rate Severity)."
            ],
            "severity": (
                "Derived from the higher of the cross-category or unit-rate repetition severity curves. Filtered against the severity floor."
            ),
            "safeguards": (
                "Single-category identical costs (e.g., 5 identical road segments) do not trigger cross-category alerts. Unit-rate checks strictly require high-confidence text extraction."
            ),
            "workflow": [
                "Group works by MP + Cost; flag groups with >= 5 works and > 1 category.",
                "Extract benchmark type, unit type, and quantity (filtering for high confidence).",
                "Compute unit rate (Cost ÷ Quantity) and round to nearest Rs. 100.",
                "Group by MP + Category + Rounded Unit Rate; flag groups with >= 5 works.",
                "Evaluate each work for cross-category clone and cloned unit rate; select maximum severity.",
                "Persist anomaly record with repeat count, rounded unit rate, and category count."
            ],
            "example": {
                "desc": "An MP has 6 projects budgeted at exactly Rs. 10,00,000:",
                "scenario": "1 Road, 1 School, 1 Water project, 1 Lighting project, 1 Community Hall, and 1 Drainage work.\nBecause there are >= 5 works with identical cost across > 1 category, Signal 1 triggers with 0.75 Severity for possible boilerplate budgeting."
            }
        }
    ]

    for det in detectors:
        # Heading 1
        h1 = doc.add_heading(level=1)
        r_h1 = h1.add_run(f"{det['id']}: {det['name']}")
        r_h1.font.name = 'Arial'
        r_h1.font.size = Pt(14)
        r_h1.font.bold = True
        r_h1.font.color.rgb = RGBColor(15, 23, 42)
        h1.paragraph_format.space_before = Pt(14)
        h1.paragraph_format.space_after = Pt(2)

        # Subtitle / Domain
        p_sub = doc.add_paragraph()
        r_sub_lbl = p_sub.add_run("Methodology: ")
        r_sub_lbl.bold = True
        r_sub_lbl.font.color.rgb = RGBColor(71, 85, 105)
        r_sub_val = p_sub.add_run(det['subtitle'])
        r_sub_val.font.color.rgb = RGBColor(37, 99, 235)
        p_sub.paragraph_format.space_after = Pt(4)

        # 1. Plain English Purpose
        p_purp = doc.add_paragraph()
        r_purp_lbl = p_purp.add_run("1. What is this detector in plain English?\n")
        r_purp_lbl.bold = True
        r_purp_lbl.font.size = Pt(11.5)
        r_purp_lbl.font.color.rgb = RGBColor(30, 58, 138)
        p_purp.add_run(det['plain_summary'])
        p_purp.paragraph_format.space_after = Pt(4)

        # 2. What it checks (Inputs)
        p_inp = doc.add_paragraph()
        r_inp_lbl = p_inp.add_run("2. What data and features does it check?")
        r_inp_lbl.bold = True
        r_inp_lbl.font.size = Pt(11.5)
        r_inp_lbl.font.color.rgb = RGBColor(30, 58, 138)
        p_inp.paragraph_format.space_after = Pt(2)

        for inp in det['inputs']:
            p_bullet = doc.add_paragraph(style='List Bullet')
            p_bullet.paragraph_format.space_after = Pt(2)
            parts = inp.split(":", 1)
            if len(parts) == 2:
                r_k = p_bullet.add_run(parts[0] + ":")
                r_k.bold = True
                p_bullet.add_run(parts[1])
            else:
                p_bullet.add_run(inp)

        # 3. Detection Rules & Thresholds
        p_rul = doc.add_paragraph()
        r_rul_lbl = p_rul.add_run("3. How does it detect issues? (Rules & Thresholds)")
        r_rul_lbl.bold = True
        r_rul_lbl.font.size = Pt(11.5)
        r_rul_lbl.font.color.rgb = RGBColor(30, 58, 138)
        p_rul.paragraph_format.space_before = Pt(4)
        p_rul.paragraph_format.space_after = Pt(2)

        for rule in det['rules']:
            p_bullet = doc.add_paragraph(style='List Bullet')
            p_bullet.paragraph_format.space_after = Pt(2)
            parts = rule.split(":", 1)
            if len(parts) == 2:
                r_k = p_bullet.add_run(parts[0] + ":")
                r_k.bold = True
                p_bullet.add_run(parts[1])
            else:
                p_bullet.add_run(rule)

        # 4. Severity & Scoring
        p_sev = doc.add_paragraph()
        r_sev_lbl = p_sev.add_run("4. Severity & Alert Levels\n")
        r_sev_lbl.bold = True
        r_sev_lbl.font.size = Pt(11.5)
        r_sev_lbl.font.color.rgb = RGBColor(30, 58, 138)
        p_sev.add_run(det['severity'])
        p_sev.paragraph_format.space_before = Pt(4)
        p_sev.paragraph_format.space_after = Pt(4)

        # 5. False-Positive Safeguards
        p_safe = doc.add_paragraph()
        r_safe_lbl = p_safe.add_run("5. False-Positive Safeguards (When is it NOT fraud?)\n")
        r_safe_lbl.bold = True
        r_safe_lbl.font.size = Pt(11.5)
        r_safe_lbl.font.color.rgb = RGBColor(30, 58, 138)
        p_safe.add_run(det['safeguards'])
        p_safe.paragraph_format.space_after = Pt(4)

        # 6. Step-by-Step Flow
        p_flow = doc.add_paragraph()
        r_flow_lbl = p_flow.add_run("6. Step-by-Step Algorithm Flow")
        r_flow_lbl.bold = True
        r_flow_lbl.font.size = Pt(11.5)
        r_flow_lbl.font.color.rgb = RGBColor(30, 58, 138)
        p_flow.paragraph_format.space_before = Pt(4)
        p_flow.paragraph_format.space_after = Pt(2)

        for step in det['workflow']:
            p_num = doc.add_paragraph(style='List Number')
            p_num.paragraph_format.space_after = Pt(2)
            p_num.add_run(step)

        # 7. Real-World Example
        p_ex = doc.add_paragraph()
        r_ex_lbl = p_ex.add_run("7. Simple Real-World Example\n")
        r_ex_lbl.bold = True
        r_ex_lbl.font.size = Pt(11.5)
        r_ex_lbl.font.color.rgb = RGBColor(30, 58, 138)
        r_ex_desc = p_ex.add_run(det['example']['desc'] + "\n")
        r_ex_desc.font.italic = True
        r_ex_scen = p_ex.add_run(det['example']['scenario'])
        p_ex.paragraph_format.space_before = Pt(4)
        p_ex.paragraph_format.space_after = Pt(14)

    # Save to file
    doc.save(output_path)
    print(f"Simple Word document successfully created at: {output_path}")

if __name__ == "__main__":
    downloads_path = "/Users/suvendu/Downloads/SATARK_MPLADS_15_Detectors_Simple_Guide.docx"
    spec_path = "/Users/suvendu/Downloads/SATARK_MPLADS_15_Detectors_Specification.docx"
    project_path = "/Users/suvendu/Downloads/SIH-DATA/SATARK_MPLADS_15_Detectors_Specification.docx"

    generate_simple_docx(downloads_path)
    generate_simple_docx(spec_path)
    generate_simple_docx(project_path)
